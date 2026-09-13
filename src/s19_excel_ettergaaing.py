"""Steg 19: Excel-arbeidsbok for å ettergå beregningene med levende formler.

Skriver output/NRK_etterproving_beregninger.xlsx. Alt som kan regnes i Excel, regnes med formler
fra rådataarkene (blå = kildedata, svart = formel i samme ark, grønn = referanse til annet ark):

  Panel + Kommunesnitt + M1   NRK-lik FE-modell (M1) og placebo (M5) som within-transformasjon
                              (variabel minus kommunesnitt) og LINEST. Gir samme koeffisienter som
                              PanelOLS med kommune-FE. Standardfeilene fra LINEST er IKKE klustret.
  DiD_data + DiD_kontrast     Hoveddesignet som gjennomsiktig kontrast: gap = log forbruk per måler
                              (bestilt) − (ikke bestilt) i samme område, klasse og måned; trekk fra
                              gapets snitt for samme kalendermåned i førperioden; snitt over de sju
                              postmånedene. Uvektet snitt er identisk med uvektet D7 (kontrolleres).
  DiD_placebo                 Samme kontrast med innføring okt. 2024 på data før okt. 2025.
  Dose                        Andel av sent-kohorten med Norgespris per måned fra daglige tellinger.
  Forbruk                     Forbruk jan.–jun. per år, endring, dusjår og vann.
  Aggregering                 Forbruksandel bestillere × effekt.
  Resultater                  Excel-formel mot Python-kontroll (regnet her) og mot pipelinens tall.

Hvorfor within + LINEST: Excel kan ikke estimere 197 kommunedummyer direkte, men FE-estimatoren er
algebraisk lik OLS på variabler fratrukket kommunesnittet. Det gjør hele M1-kjeden etterprøvbar.
Kjøres etter s06–s18. Filen recalculeres ved åpning (fullCalcOnLoad); verdier vises ikke før da.
"""
import json
import numpy as np
import pandas as pd
import statsmodels.formula.api as smf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from config import PROC, OUT, ROOT

VERIF = ROOT / "verifisering"
XLSX = OUT / "NRK_etterproving_beregninger.xlsx"
POST = "2025-10"

BLUE, BLACK, GREEN = Font(color="0000FF"), Font(color="000000"), Font(color="008000")
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="E7EEF5")
NOTE = Font(italic=True, color="595959")


def sheet(wb, name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, row, cols, widths=None):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = BOLD; cell.fill = HEAD; cell.alignment = Alignment(wrap_text=True, vertical="top")
        if widths:
            ws.column_dimensions[L(j)].width = widths[j - 1]


def put(ws, row, col, value, font=BLACK, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    if fmt:
        c.number_format = fmt
    return c


# ----------------------------------------------------------------------------- data
did = pd.read_parquet(PROC / "did_month.parquet").sort_values(["pa", "eac", "status", "maned"]).reset_index(drop=True)
did["mnd"] = did.maned.str[5:]
pan = pd.read_parquet(PROC / "panel.parquet").sort_values(["maned", "knr"]).reset_index(drop=True)
pan["innt_1000"] = pan.median_innt_etter_skatt / 1000
n_pre = int((pan.maned < POST).sum())
assert (pan.maned.iloc[:n_pre] < POST).all() and (pan.maned.iloc[n_pre:] >= POST).all()
daily = pd.read_parquet(PROC / "elhub_np_mba_daily.parquet")
daily = daily[(daily.FORBRUKSGRUPPE == "Husholdning") & daily.PRISOMRADE.isin(["NO1", "NO2", "NO5"])].copy()
daily["dag"] = daily.DATO.astype(str).str[:10]
daily = daily[(daily.dag >= "2025-10-01") & (daily.dag <= "2026-04-30")].sort_values(["PRISOMRADE", "dag"]).reset_index(drop=True)
cons = pd.read_parquet(PROC / "elhub_cons_mba_month.parquet")
cons = cons[cons.PRISOMRADE.isin(["NO1", "NO2", "NO5"]) & cons.FORBRUKSGRUPPE.isin(["Husholdning", "Hytter og fritidseiendommer"])]
cons = cons.sort_values(["PRISOMRADE", "FORBRUKSGRUPPE", "maned"]).reset_index(drop=True)
cons["FORBRUKSGRUPPE"] = cons.FORBRUKSGRUPPE.astype(str)

# ----------------------------------------------------------------------------- Python-kontroll (samme kjede som formlene)
def within_ols(df, y, xs):
    d = df.copy()
    for c in [y] + xs:
        d[c + "_w"] = d[c] - d.groupby("knr")[c].transform("mean")
    X = np.column_stack([np.ones(len(d))] + [d[c + "_w"] for c in xs])
    b, *_ = np.linalg.lstsq(X, d[y + "_w"], rcond=None)
    res = d[y + "_w"] - X @ b
    r2 = 1 - (res ** 2).sum() / ((d[y + "_w"] - d[y + "_w"].mean()) ** 2).sum()
    return dict(zip(xs, b[1:])), r2

XS = ["egd", "spot_ore_kwh", "innt_1000", "snitt_bruksareal_m2", "post_np"]
m1_b, m1_r2 = within_ols(pan, "log_kwh", XS)
pre = pan[pan.maned < POST].copy(); pre["post_np"] = (pre.maned >= "2024-10").astype(int)
m5_b, m5_r2 = within_ols(pre, "log_kwh", XS)

def kontrast(df, post_from):
    k = df[df.status == "Ikke bestilt"][["pa", "eac", "maned", "log_y"]].rename(columns={"log_y": "y_k"})
    p = df[df.status != "Ikke bestilt"].merge(k, on=["pa", "eac", "maned"])
    p["gap"] = p.log_y - p.y_k; p["post"] = (p.maned >= post_from).astype(int)
    g = p[p.post == 0].groupby(["pa", "eac", "status", "mnd"]).gap.mean().rename("gap_for").reset_index()
    e = p[p.post == 1].merge(g, on=["pa", "eac", "status", "mnd"])
    e["avvik"] = e.gap - e.gap_for
    return e.avvik.mean(), np.average(e.avvik, weights=e.n_mp), e.groupby("maned").avvik.mean()

k_uv, k_v, k_mnd = kontrast(did, POST)
pl_uv, pl_v, _ = kontrast(did[did.maned < POST], "2024-10")
did["bestilt_post"] = (did.status != "Ikke bestilt").astype(int) * (did.maned >= POST).astype(int)
did["pa_t_eac"] = did.pa + "|" + did.maned + "|" + did.eac; did["bestilt"] = (did.status != "Ikke bestilt").astype(int)
ols_d7 = smf.ols("log_y ~ C(g) + C(pa_t_eac) + bestilt:C(mnd):C(eac) + bestilt_post", data=did).fit().params["bestilt_post"]
print(f"Kontrast uvektet {k_uv:.6f} | uvektet OLS-D7 {ols_d7:.6f} | avvik {k_uv - ols_d7:.2e}")
print(f"Kontrast n_mp-vektet {k_v:.6f} | placebo uvektet {pl_uv:.6f}")

dose_py = {}
for pa, g in daily.groupby("PRISOMRADE"):
    g = g.set_index("dag"); n0, nT = g.loc["2025-10-01", "ANTALL_NORGESPRIS"], g.loc["2026-04-30", "ANTALL_NORGESPRIS"]
    d = ((g.ANTALL_NORGESPRIS - n0) / (nT - n0)).clip(0, 1); dose_py[pa] = d.groupby(g.index.str[:7]).mean()
cons["aar"] = cons.maned.str[:4].astype(int); cons["mnd"] = cons.maned.str[5:].astype(int)
h1 = cons[cons.mnd <= 6].groupby(["aar", "FORBRUKSGRUPPE"]).kwh.sum().unstack() / 1e6
h1["sum"] = h1.sum(axis=1)
post_kwh = did[did.maned >= POST].groupby("bestilt").kwh.sum(); andel_py = post_kwh[1] / post_kwh.sum()
print(f"Forbruksandel bestillere {andel_py:.4f} | forbruk 2026 H1 {h1.loc[2026, 'sum']:.0f} GWh")

# ----------------------------------------------------------------------------- pipeline-tall
def csv_val(path, **crit):
    try:
        t = pd.read_csv(path)
        for k, v in crit.items():
            t = t[t[k].astype(str).str.startswith(str(v))] if k == "modell" else t[t[k] == v]
        return t
    except Exception:
        return pd.DataFrame()

fe = pd.read_csv(OUT / "tab_fe_resultater.csv").rename(columns={"Unnamed: 0": "var"})
fe_m1 = fe[fe.modell == "M1 NRK-lik"].set_index("var").koef
fe_m5 = fe[fe.modell.str.startswith("M5 placebo")].set_index("var").koef
mm = pd.read_csv(OUT / "tab_fe_modellmaal.csv").set_index("modell").R2_within
d7 = csv_val(OUT / "tab_did_eac_sesong.csv", modell="D7 område", term="bestilt_post")
d7p = csv_val(OUT / "tab_did_eac_sesong.csv", modell="D7 placebo", term="bestilt_post")
wcb = csv_val(OUT / "tab_wcb.csv", modell="D7 område", term="bestilt_post")
vk = pd.read_csv(VERIF / "vekting_og_klustring.csv").set_index("modell").koef if (VERIF / "vekting_og_klustring.csv").exists() else pd.Series(dtype=float)
p1 = pd.read_csv(OUT / "tab_p1_forbruk_jan_jun.csv").set_index("aar")
try:
    dsj = json.loads((VERIF / "dose_sensitivitet.json").read_text(encoding="utf-8")); andel_pipe = dsj["andel_bestilt_av_postforbruk"]
except Exception:
    andel_pipe = None

PIPE = [  # (etikett, verdi, kilde)
    ("M1 egd", fe_m1.get("egd"), "output/tab_fe_resultater.csv, M1 NRK-lik"),
    ("M1 spot", fe_m1.get("spot_ore_kwh"), "output/tab_fe_resultater.csv"),
    ("M1 innt_1000", fe_m1.get("innt_1000"), "output/tab_fe_resultater.csv"),
    ("M1 bruksareal", fe_m1.get("snitt_bruksareal_m2"), "output/tab_fe_resultater.csv"),
    ("M1 post_np", fe_m1.get("post_np"), "output/tab_fe_resultater.csv"),
    ("M1 R2 within", mm.get("M1 NRK-lik"), "output/tab_fe_modellmaal.csv"),
    ("M5 post_np", fe_m5.get("post_np"), "output/tab_fe_resultater.csv, M5 placebo"),
    ("M5 R2 within", mm.get("M5 placebo 2024-10 (data ≤ 2025-09)"), "output/tab_fe_modellmaal.csv"),
    ("D7 bestilt_post (WLS, n_mp)", float(d7.koef.iloc[0]) if len(d7) else None, "output/tab_did_eac_sesong.csv"),
    ("D7 klustret SE", float(d7.se.iloc[0]) if len(d7) else None, "output/tab_did_eac_sesong.csv"),
    ("D7 uvektet OLS", float(vk.get("OLS_g")) if "OLS_g" in vk.index else None, "verifisering/vekting_og_klustring.csv"),
    ("D7 placebo (WLS)", float(d7p.koef.iloc[0]) if len(d7p) else None, "output/tab_did_eac_sesong.csv"),
    ("WCB KI lav (log)", float(wcb.ki_wcb_lav.iloc[0]) if len(wcb) else None, "output/tab_wcb.csv"),
    ("WCB KI høy (log)", float(wcb.ki_wcb_hoy.iloc[0]) if len(wcb) else None, "output/tab_wcb.csv"),
    ("Forbruk H1 2026, hush+hytter GWh", float(p1.loc[2026, "Hus+hytter"]), "output/tab_p1_forbruk_jan_jun.csv"),
    ("Forbruk H1 2025, hush+hytter GWh", float(p1.loc[2025, "Hus+hytter"]), "output/tab_p1_forbruk_jan_jun.csv"),
    ("Forbruksandel bestillere (post)", andel_pipe, "verifisering/dose_sensitivitet.json"),
]

# ----------------------------------------------------------------------------- arbeidsbok
wb = Workbook(); wb.remove(wb.active)
wb.calculation.fullCalcOnLoad = True

# --- Les meg
ws = sheet(wb, "LesMeg"); ws.column_dimensions["A"].width = 120
lines = [
    ("Etterprøving av NRK «Prisen for billig strøm» (27.08.2026): beregningene med levende formler", BOLD),
    ("Laget av src/s19_excel_ettergaaing.py. Alle tall er til godkjenning og er ikke NRKs tall. Filen recalculeres når den åpnes i Excel; det tar noen sekunder (LINEST på 10 047 rader).", NOTE),
    ("", None),
    ("Fargekode: blå = kildedata/input, svart = formel i samme ark, grønn = referanse til annet ark.", None),
    ("", None),
    ("Ark og hva de viser:", BOLD),
    ("Resultater – hver størrelse regnet tre ganger: Excel-formel, Python (samme kjede, i dette skriptet) og pipelinen (statsmodels/linearmodels). Avvikskolonner viser differansen.", None),
    ("Panel – kommune × måned, 197 kommuner, feb. 2022–apr. 2026 (Elhub, Frost, hvakosterstrommen, SSB). Kolonner K–N er avledet, O–Z er variabler fratrukket kommunesnittet (within-transformasjon), for M1 (O–T) og M5 (U–Z).", None),
    ("Kommunesnitt – snitt per kommune som trekkes fra i Panel; for M5 bare over måneder før okt. 2025.", None),
    ("M1 – NRK-lik modell og placebo M5 med LINEST på de within-transformerte kolonnene. Koeffisientene er identiske med en modell med faste kommuneeffekter. Standardfeilene fra LINEST er ikke klustret og skal ikke brukes til slutning; pipelinen bruker klustrede og Driscoll-Kraay-standardfeil.", None),
    ("DiD_data – Elhubs uttrekk etter bestillingsstatus: 27 grupper (3 prisområder × 3 statuser × 3 forbruksklasser) × 31 måneder. kWh og antall målere er kildedata; kWh per måler og logaritmen er formler.", None),
    ("DiD_kontrast – hoveddesignet regnet gjennomsiktig: gap = log forbruk per måler for bestilt minus ikke bestilt i samme område, klasse og måned; gap_for = snitt av gapet for samme kalendermåned i de to årene før ordningen; avvik = gap − gap_for for månedene fra okt. 2025. Snittet av avvikene er effekten. Det uvektede snittet er identisk med den uvektede regresjonsmodellen D7 (se Resultater). Pipelinens hovedtall vekter med antall målere i selve regresjonen (0,0299); det vektede snittet her (SUMPRODUCT) er en tilnærming og avviker litt fra dette.", None),
    ("DiD_placebo – samme kontrast med innføring satt til okt. 2024 og bare data før okt. 2025. Førperioden er da ett år per kalendermåned.", None),
    ("Dose – andel av kohorten «bestilt sent» med aktiv Norgespris per måned, fra Elhubs daglige tellinger av alle husholdningsmålere (forutsetter at kohorten følger totalen).", None),
    ("Forbruk – husholdninger og hytter i NO1, NO2, NO5, januar–juni per år; endring, dusjår (én dusj kontinuerlig i ett år = 1,4 kWh per 6 min × 8 760 timer) og vann (1 429 liter per kWh, NRKs tall).", None),
    ("Aggregering – bestillernes andel av husholdningsforbruket etter okt. 2025 × effekt for bestillere = effekt for alle husholdninger, forutsatt at ikke-bestillerne er upåvirket.", None),
    ("Pipeline – tall hentet fra pipelinens CSV-filer ved bygging (blå), med kildefil.", None),
    ("", None),
    ("Kilder: Elhub (data.elhub.no, CC BY 4.0), Meteorologisk institutt Frost (CC BY 4.0), hvakosterstrommen.no, SSB tabell 06944 og 06513 (CC BY 4.0). Datasettnavn står i src/config.py. Kode: github.com/jakobes1999-boop/NRK.", None),
    ("Forbehold: kommune→prisområde uten offisiell kilde (28 av 29 kontrollerte bekreftet mot VG); Elhub-gruppen «Privat» i kommunefilen er husholdning og hytte samlet; inntekt framskrevet etter 2024; DiD-kontrasten er mot strømstøtte, ikke markedspris; seleksjon inn i bestilling, se notat/METODENOTAT_til_NRK_2026-09-11.md.", None),
    ("Dato: 13. september 2026.", None),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t); c.alignment = Alignment(wrap_text=True, vertical="top")
    if f: c.font = f

# --- Pipeline
wp = sheet(wb, "Pipeline"); header(wp, 1, ["Størrelse", "Verdi", "Kilde"], [36, 16, 60])
pipe_row = {}
for i, (lab, val, src) in enumerate(PIPE, 2):
    put(wp, i, 1, lab); put(wp, i, 2, val if val is not None else "[fylles inn]", BLUE, "0.000000"); put(wp, i, 3, src, NOTE); pipe_row[lab] = i

# --- Panel
wpan = sheet(wb, "Panel")
cols = ["knr", "kommune", "maned", "kwh", "n_mp", "egd", "spot_ore_kwh", "innt_kr", "snitt_bruksareal_m2", "post_np",
        "log_kwh", "innt_1000", "kwh_per_mp", "post_2024 (M5)",
        "w_log_kwh", "w_egd", "w_spot", "w_innt_1000", "w_areal", "w_post",
        "w5_log_kwh", "w5_egd", "w5_spot", "w5_innt_1000", "w5_areal", "w5_post_2024"]
header(wpan, 1, cols, [7, 16, 9, 14, 10, 8, 8, 10, 8, 7, 9, 9, 9, 9] + [10] * 12)
wpan.freeze_panes = "A2"
N = len(pan); last = N + 1
src_cols = {"K": "F", "F": "F", "G": "G", "L": "L", "I": "I", "J": "J"}  # kolonner som demeanes (M1)
for r, row in enumerate(pan.itertuples(index=False), 2):
    vals = [row.knr, row.kommune, row.maned, float(row.kwh), float(row.n_mp), float(row.egd), float(row.spot_ore_kwh),
            float(row.median_innt_etter_skatt), float(row.snitt_bruksareal_m2), int(row.post_np)]
    for j, v in enumerate(vals, 1):
        put(wpan, r, j, v, BLUE, "0.00" if j in (6, 7, 9) else None)
    put(wpan, r, 11, f"=LN(D{r})", fmt="0.0000"); put(wpan, r, 12, f"=H{r}/1000", fmt="0.0"); put(wpan, r, 13, f"=D{r}/E{r}", fmt="0")
    put(wpan, r, 14, f'=IF(C{r}<"{POST}",IF(C{r}>="2024-10",1,0),0)')
    for j, (src, ks) in enumerate(zip(["K", "F", "G", "L", "I", "J"], ["B", "C", "D", "E", "F", "G"]), 15):
        put(wpan, r, j, f"={src}{r}-INDEX(Kommunesnitt!${ks}:${ks},MATCH($A{r},Kommunesnitt!$A:$A,0))", GREEN, "0.0000")
    for j, (src, ks) in enumerate(zip(["K", "F", "G", "L", "I", "N"], ["H", "I", "J", "K", "L", "M"]), 21):
        put(wpan, r, j, f"={src}{r}-INDEX(Kommunesnitt!${ks}:${ks},MATCH($A{r},Kommunesnitt!$A:$A,0))", GREEN, "0.0000")

# --- Kommunesnitt
wk = sheet(wb, "Kommunesnitt")
header(wk, 1, ["knr", "M1 log_kwh", "M1 egd", "M1 spot", "M1 innt_1000", "M1 areal", "M1 post_np",
               "M5 log_kwh", "M5 egd", "M5 spot", "M5 innt_1000", "M5 areal", "M5 post_2024"], [8] + [12] * 12)
put(wk, 1, 15, "M5-snitt er over måneder før okt. 2025 (samme utvalg som placeboen).", NOTE)
for r, knr in enumerate(sorted(pan.knr.unique()), 2):
    put(wk, r, 1, knr, BLUE)
    for j, src in enumerate(["K", "F", "G", "L", "I", "J"], 2):
        put(wk, r, j, f"=AVERAGEIFS(Panel!${src}:${src},Panel!$A:$A,$A{r})", GREEN, "0.0000")
    for j, src in enumerate(["K", "F", "G", "L", "I", "N"], 8):
        put(wk, r, j, f'=AVERAGEIFS(Panel!${src}:${src},Panel!$A:$A,$A{r},Panel!$C:$C,"<{POST}")', GREEN, "0.0000")

# --- M1
wm = sheet(wb, "M1"); wm.column_dimensions["A"].width = 30
for c in "BCDEF": wm.column_dimensions[c].width = 16
put(wm, 1, 1, "M1: NRK-lik modell. log(kWh) på within-transformerte variabler (Panel O–T), LINEST.", BOLD)
put(wm, 2, 1, "Koeffisientene er lik en modell med faste kommuneeffekter. LINEST-standardfeil er ikke klustret; se Pipeline for klustrede/Driscoll-Kraay.", NOTE)
header(wm, 4, ["Variabel", "Koeffisient (Excel)", "SE (LINEST, ukorrigert)", "Python-kontroll", "Pipeline (PanelOLS)", "Avvik Excel−pipeline"])
names = ["egd", "spot_ore_kwh", "innt_1000", "snitt_bruksareal_m2", "post_np"]
lin1 = f"LINEST(Panel!$O$2:$O${last},Panel!$P$2:$T${last},TRUE,TRUE)"
lin5 = f"LINEST(Panel!$U$2:$U${n_pre + 1},Panel!$V$2:$Z${n_pre + 1},TRUE,TRUE)"
m1_row = {}
for i, nm in enumerate(names):
    r = 5 + i; idx = len(names) - i  # LINEST returnerer siste x-kolonne først
    put(wm, r, 1, nm); put(wm, r, 2, f"=INDEX({lin1},1,{idx})", GREEN, "0.000000"); put(wm, r, 3, f"=INDEX({lin1},2,{idx})", GREEN, "0.000000")
    put(wm, r, 4, float(m1_b[nm]), BLUE, "0.000000")
    put(wm, r, 5, f"=Pipeline!B{pipe_row['M1 ' + {'egd':'egd','spot_ore_kwh':'spot','innt_1000':'innt_1000','snitt_bruksareal_m2':'bruksareal','post_np':'post_np'}[nm]]}", GREEN, "0.000000")
    put(wm, r, 6, f"=B{r}-E{r}", fmt="0.000000"); m1_row[nm] = r
r = 10; put(wm, r, 1, "R² (within)"); put(wm, r, 2, f"=INDEX({lin1},3,1)", GREEN, "0.0000"); put(wm, r, 4, float(m1_r2), BLUE, "0.0000")
put(wm, r, 5, f"=Pipeline!B{pipe_row['M1 R2 within']}", GREEN, "0.0000"); put(wm, r, 6, f"=B{r}-E{r}", fmt="0.0000")
r = 11; put(wm, r, 1, "Norgespris-effekt, prosent"); put(wm, r, 2, f"=(EXP(B{m1_row['post_np']})-1)*100", fmt="0.00")
put(wm, r, 4, float((np.exp(m1_b["post_np"]) - 1) * 100), BLUE, "0.00")
put(wm, 14, 1, f"M5: placebo. Innføring okt. 2024, kun rader før okt. 2025 (Panel rad 2–{n_pre + 1}), within-transformerte kolonner U–Z.", BOLD)
header(wm, 16, ["Variabel", "Koeffisient (Excel)", "SE (LINEST, ukorrigert)", "Python-kontroll", "Pipeline (PanelOLS)", "Avvik Excel−pipeline"])
for i, nm in enumerate(names):
    r = 17 + i; idx = len(names) - i
    put(wm, r, 1, nm if nm != "post_np" else "post_2024 (placebo)"); put(wm, r, 2, f"=INDEX({lin5},1,{idx})", GREEN, "0.000000"); put(wm, r, 3, f"=INDEX({lin5},2,{idx})", GREEN, "0.000000")
    put(wm, r, 4, float(m5_b[nm]), BLUE, "0.000000")
    if nm == "post_np":
        put(wm, r, 5, f"=Pipeline!B{pipe_row['M5 post_np']}", GREEN, "0.000000"); put(wm, r, 6, f"=B{r}-E{r}", fmt="0.000000"); m5_post_row = r
r = 22; put(wm, r, 1, "R² (within)"); put(wm, r, 2, f"=INDEX({lin5},3,1)", GREEN, "0.0000"); put(wm, r, 4, float(m5_r2), BLUE, "0.0000")
put(wm, r, 5, f"=Pipeline!B{pipe_row['M5 R2 within']}", GREEN, "0.0000"); put(wm, r, 6, f"=B{r}-E{r}", fmt="0.0000")
r = 23; put(wm, r, 1, "Placebo-«effekt», prosent"); put(wm, r, 2, f"=(EXP(B{m5_post_row})-1)*100", fmt="0.00"); put(wm, r, 4, float((np.exp(m5_b["post_np"]) - 1) * 100), BLUE, "0.00")

# --- DiD_data
wd = sheet(wb, "DiD_data")
header(wd, 1, ["pa", "status", "eac", "maned", "mnd", "post", "kwh", "n_mp", "kwh_per_mp", "log_y", "bestilt"], [6, 14, 18, 9, 6, 6, 14, 10, 10, 9, 7])
wd.freeze_panes = "A2"
for r, row in enumerate(did.itertuples(index=False), 2):
    put(wd, r, 1, row.pa, BLUE); put(wd, r, 2, row.status, BLUE); put(wd, r, 3, row.eac, BLUE); put(wd, r, 4, row.maned, BLUE)
    put(wd, r, 5, f"=RIGHT(D{r},2)"); put(wd, r, 6, f'=IF(D{r}>="{POST}",1,0)')
    put(wd, r, 7, float(row.kwh), BLUE, "#,##0"); put(wd, r, 8, float(row.n_mp), BLUE, "#,##0")
    put(wd, r, 9, f"=G{r}/H{r}", fmt="0.0"); put(wd, r, 10, f"=LN(I{r})", fmt="0.0000"); put(wd, r, 11, f'=IF(B{r}="Ikke bestilt",0,1)')
dlast = len(did) + 1

def kontrast_ark(name, df, post_from, title):
    w = sheet(wb, name)
    header(w, 1, ["pa", "eac", "status", "maned", "mnd", "post", "log_y bestilt", "log_y ikke bestilt", "gap", "gap_for (snitt før, samme kalendermåned)", "avvik (post)", "n_mp bestilt"],
           [6, 18, 14, 9, 6, 6, 12, 12, 10, 14, 12, 12])
    w.freeze_panes = "A2"
    rows = df[df.status != "Ikke bestilt"][["pa", "eac", "status", "maned"]].drop_duplicates().sort_values(["pa", "eac", "status", "maned"]).reset_index(drop=True)
    for r, row in enumerate(rows.itertuples(index=False), 2):
        put(w, r, 1, row.pa, BLUE); put(w, r, 2, row.eac, BLUE); put(w, r, 3, row.status, BLUE); put(w, r, 4, row.maned, BLUE)
        put(w, r, 5, f"=RIGHT(D{r},2)"); put(w, r, 6, f'=IF(D{r}>="{post_from}",1,0)')
        put(w, r, 7, f"=SUMIFS(DiD_data!$J:$J,DiD_data!$A:$A,$A{r},DiD_data!$C:$C,$B{r},DiD_data!$B:$B,$C{r},DiD_data!$D:$D,$D{r})", GREEN, "0.0000")
        put(w, r, 8, f'=SUMIFS(DiD_data!$J:$J,DiD_data!$A:$A,$A{r},DiD_data!$C:$C,$B{r},DiD_data!$B:$B,"Ikke bestilt",DiD_data!$D:$D,$D{r})', GREEN, "0.0000")
        put(w, r, 9, f"=G{r}-H{r}", fmt="0.0000")
        put(w, r, 10, f"=AVERAGEIFS($I:$I,$A:$A,$A{r},$B:$B,$B{r},$C:$C,$C{r},$E:$E,$E{r},$F:$F,0)", fmt="0.0000")
        put(w, r, 11, f"=IF($F{r}=1,$I{r}-$J{r},0)", fmt="0.0000")
        put(w, r, 12, f"=SUMIFS(DiD_data!$H:$H,DiD_data!$A:$A,$A{r},DiD_data!$C:$C,$B{r},DiD_data!$B:$B,$C{r},DiD_data!$D:$D,$D{r})", GREEN, "#,##0")
    n = len(rows) + 1
    put(w, 1, 14, title, BOLD); w.column_dimensions["N"].width = 34; w.column_dimensions["O"].width = 14
    put(w, 2, 14, "Uvektet snitt av avvik (post), log"); put(w, 2, 15, f"=SUMIF($F$2:$F${n},1,$K$2:$K${n})/COUNTIF($F$2:$F${n},1)", fmt="0.000000")
    put(w, 3, 14, "Tilsvarende prosent"); put(w, 3, 15, "=(EXP(O2)-1)*100", fmt="0.00")
    put(w, 4, 14, "Vektet med n_mp bestilt (tilnærming)"); put(w, 4, 15, f"=SUMPRODUCT($F$2:$F${n},$L$2:$L${n},$K$2:$K${n})/SUMPRODUCT($F$2:$F${n},$L$2:$L${n})", fmt="0.000000")
    put(w, 5, 14, "Antall postobservasjoner"); put(w, 5, 15, f"=COUNTIF($F$2:$F${n},1)")
    put(w, 7, 14, "Per måned (uvektet snitt av avvik, prosent)", BOLD)
    post_m = sorted(rows[rows.maned >= post_from].maned.unique())
    for i, m in enumerate(post_m):
        put(w, 8 + i, 14, m, BLUE); put(w, 8 + i, 15, f'=(EXP(AVERAGEIFS($K$2:$K${n},$F$2:$F${n},1,$D$2:$D${n},N{8 + i}))-1)*100', fmt="0.00")
    return w

kontrast_ark("DiD_kontrast", did, POST, "Hoveddesign: effekt = snitt av avvik fra eget førmønster")
kontrast_ark("DiD_placebo", did[did.maned < POST], "2024-10", "Placebo: innføring okt. 2024, data før okt. 2025")

# --- Dose
wdo = sheet(wb, "Dose")
header(wdo, 1, ["dag", "pa", "antall_norgespris", "antall målere", "maned", "N 1. okt.", "N 30. apr.", "dose"], [11, 6, 14, 12, 9, 12, 12, 8])
wdo.freeze_panes = "A2"
for r, row in enumerate(daily.itertuples(index=False), 2):
    put(wdo, r, 1, row.dag, BLUE); put(wdo, r, 2, row.PRISOMRADE, BLUE); put(wdo, r, 3, float(row.ANTALL_NORGESPRIS), BLUE, "#,##0"); put(wdo, r, 4, float(row.ANTALL), BLUE, "#,##0")
    put(wdo, r, 5, f"=LEFT(A{r},7)")
    put(wdo, r, 6, f'=SUMIFS($C:$C,$B:$B,B{r},$A:$A,"2025-10-01")', fmt="#,##0"); put(wdo, r, 7, f'=SUMIFS($C:$C,$B:$B,B{r},$A:$A,"2026-04-30")', fmt="#,##0")
    put(wdo, r, 8, f"=MAX(0,MIN(1,(C{r}-F{r})/(G{r}-F{r})))", fmt="0.000")
dn = len(daily) + 1
put(wdo, 1, 10, "Dose per måned = snitt av daglig dose (Excel)", BOLD); wdo.column_dimensions["J"].width = 10
months = ["2025-10", "2025-11", "2025-12", "2026-01", "2026-02", "2026-03", "2026-04"]
for j, pa in enumerate(["NO1", "NO2", "NO5"]):
    put(wdo, 2, 11 + j, pa, BOLD); put(wdo, 2, 15 + j, pa + " Python", BOLD)
for i, m in enumerate(months):
    put(wdo, 3 + i, 10, m, BLUE)
    for j, pa in enumerate(["NO1", "NO2", "NO5"]):
        put(wdo, 3 + i, 11 + j, f"=AVERAGEIFS($H$2:$H${dn},$B$2:$B${dn},{L(11 + j)}$2,$E$2:$E${dn},$J{3 + i})", fmt="0.000")
        put(wdo, 3 + i, 15 + j, float(dose_py[pa].get(m, np.nan)), BLUE, "0.000")
put(wdo, 11, 10, "Dosen forutsetter at sent-kohorten følger tilstrømningen blant alle husholdningsmålere. Kohortens eget forløp er ikke publisert.", NOTE)

# --- Forbruk
wf = sheet(wb, "Forbruk")
header(wf, 1, ["pa", "gruppe", "maned", "kwh", "aar", "mnd"], [6, 26, 9, 16, 6, 5])
wf.freeze_panes = "A2"
for r, row in enumerate(cons.itertuples(index=False), 2):
    put(wf, r, 1, row.PRISOMRADE, BLUE); put(wf, r, 2, row.FORBRUKSGRUPPE, BLUE); put(wf, r, 3, row.maned, BLUE); put(wf, r, 4, float(row.kwh), BLUE, "#,##0")
    put(wf, r, 5, f"=VALUE(LEFT(C{r},4))"); put(wf, r, 6, f"=VALUE(RIGHT(C{r},2))")
fn = len(cons) + 1
put(wf, 1, 8, "Januar–juni, NO1+NO2+NO5 (Excel)", BOLD)
for c, wdt in zip("HIJKLMNOP", [8, 14, 14, 14, 14, 10, 12, 14, 16]): wf.column_dimensions[c].width = wdt
put(wf, 2, 8, "Parametre", BOLD); put(wf, 3, 8, "kWh per dusj (6 min)"); put(wf, 3, 9, 1.4, BLUE, "0.0")
put(wf, 4, 8, "Dusjer per time"); put(wf, 4, 9, 10, BLUE); put(wf, 5, 8, "Timer per år"); put(wf, 5, 9, 8760, BLUE)
put(wf, 6, 8, "kWh per dusjår"); put(wf, 6, 9, "=I3*I4*I5", fmt="#,##0"); put(wf, 7, 8, "Liter per kWh (NRK)"); put(wf, 7, 9, 1429, BLUE)
header(wf, 9, ["aar", "Husholdning GWh", "Hytter GWh", "Sum GWh", "Endring GWh", "Endring %", "Dusjår", "Vann mrd. liter", "Python-kontroll sum GWh"][0:9])
for j, t in enumerate(["aar", "Husholdning GWh", "Hytter GWh", "Sum GWh", "Endring GWh", "Endring %", "Dusjår", "Vann mrd. liter", "Python-kontroll sum GWh"]):
    c = wf.cell(row=9, column=8 + j, value=t); c.font = BOLD; c.fill = HEAD
for i, yr in enumerate(range(2021, 2027)):
    r = 10 + i
    put(wf, r, 8, yr, BLUE)
    put(wf, r, 9, f'=SUMIFS($D$2:$D${fn},$B$2:$B${fn},"Husholdning",$E$2:$E${fn},H{r},$F$2:$F${fn},"<=6")/1000000', fmt="#,##0")
    put(wf, r, 10, f'=SUMIFS($D$2:$D${fn},$B$2:$B${fn},"Hytter og fritidseiendommer",$E$2:$E${fn},H{r},$F$2:$F${fn},"<=6")/1000000', fmt="#,##0")
    put(wf, r, 11, f"=I{r}+J{r}", fmt="#,##0")
    if i > 0:
        put(wf, r, 12, f"=K{r}-K{r - 1}", fmt="#,##0"); put(wf, r, 13, f"=L{r}/K{r - 1}*100", fmt="0.0")
        put(wf, r, 14, f"=L{r}*1000000/$I$6", fmt="#,##0"); put(wf, r, 15, f"=L{r}*1000000*$I$7/1000000000", fmt="#,##0")
    put(wf, r, 16, float(h1.loc[yr, "sum"]) if yr in h1.index else None, BLUE, "#,##0")
put(wf, 17, 8, "Bare måneder til og med juni telles. Husholdning alene, endring 2026: =I15/I14−1.", NOTE)
put(wf, 18, 8, "Husholdning alene, endring % 2026"); put(wf, 18, 9, "=(I15/I14-1)*100", fmt="0.0")

# --- Aggregering
wa = sheet(wb, "Aggregering"); wa.column_dimensions["A"].width = 52; wa.column_dimensions["B"].width = 16; wa.column_dimensions["C"].width = 16
header(wa, 1, ["Størrelse", "Excel", "Python-kontroll"])
put(wa, 2, 1, "kWh bestillere, okt. 2025–apr. 2026 (DiD_data)"); put(wa, 2, 2, "=SUMIFS(DiD_data!$G:$G,DiD_data!$F:$F,1,DiD_data!$K:$K,1)", GREEN, "#,##0"); put(wa, 2, 3, float(post_kwh[1]), BLUE, "#,##0")
put(wa, 3, 1, "kWh alle, okt. 2025–apr. 2026"); put(wa, 3, 2, "=SUMIFS(DiD_data!$G:$G,DiD_data!$F:$F,1)", GREEN, "#,##0"); put(wa, 3, 3, float(post_kwh.sum()), BLUE, "#,##0")
put(wa, 4, 1, "Forbruksandel bestillere"); put(wa, 4, 2, "=B2/B3", fmt="0.000"); put(wa, 4, 3, float(andel_py), BLUE, "0.000")
put(wa, 5, 1, "Effekt for bestillere, prosent (pipeline D7, vektet)"); put(wa, 5, 2, f"=(EXP(Pipeline!B{pipe_row['D7 bestilt_post (WLS, n_mp)']})-1)*100", GREEN, "0.00")
put(wa, 6, 1, "Effekt for bestillere, prosent (Excel-kontrast, uvektet)"); put(wa, 6, 2, "=DiD_kontrast!O3", GREEN, "0.00"); put(wa, 6, 3, float((np.exp(k_uv) - 1) * 100), BLUE, "0.00")
put(wa, 7, 1, "Aggregert effekt, prosent (andel × D7)"); put(wa, 7, 2, "=B4*B5", fmt="0.00")
put(wa, 8, 1, "Aggregert effekt, prosent (andel × Excel-kontrast)"); put(wa, 8, 2, "=B4*B6", fmt="0.00"); put(wa, 8, 3, float(andel_py * (np.exp(k_uv) - 1) * 100), BLUE, "0.00")
put(wa, 9, 1, "Forbruksøkning jan.–jun. 2026, prosent (Forbruk)"); put(wa, 9, 2, "=Forbruk!M15", GREEN, "0.0")
put(wa, 10, 1, "Andel av økningen forklart av Norgespris (D7)"); put(wa, 10, 2, "=B7/B9", fmt="0.00")
put(wa, 11, 1, "Forutsetter at ikke-bestillerne er upåvirket, og at vintereffekten holder for mai–juni.", NOTE)

# --- Resultater
wr = sheet(wb, "Resultater")
header(wr, 1, ["Størrelse", "Excel-formel", "Python-kontroll", "Pipeline", "Avvik Excel−Python", "Avvik Excel−pipeline", "Kommentar"], [44, 14, 14, 14, 14, 14, 70])
R = [
    ("M1 Norgespris-koeffisient (post_np, log)", f"=M1!B{m1_row['post_np']}", float(m1_b["post_np"]), f"=Pipeline!B{pipe_row['M1 post_np']}", "Within + LINEST = FE. Skal være 0 avvik."),
    ("M1 energigradtall", f"=M1!B{m1_row['egd']}", float(m1_b["egd"]), f"=Pipeline!B{pipe_row['M1 egd']}", "NRK oppgir 0,0019."),
    ("M1 spotpris", f"=M1!B{m1_row['spot_ore_kwh']}", float(m1_b["spot_ore_kwh"]), f"=Pipeline!B{pipe_row['M1 spot']}", "NRK oppgir −0,0005."),
    ("M1 R² within", "=M1!B10", float(m1_r2), f"=Pipeline!B{pipe_row['M1 R2 within']}", "NRK oppgir 93,76 prosent."),
    ("M5 placebo-koeffisient (log)", f"=M1!B{m5_post_row}", float(m5_b["post_np"]), f"=Pipeline!B{pipe_row['M5 post_np']}", "«Effekt» av innføring okt. 2024 uten Norgespris."),
    ("DiD uvektet kontrast (log)", "=DiD_kontrast!O2", float(k_uv), f"=Pipeline!B{pipe_row['D7 uvektet OLS']}", "Identisk med uvektet regresjon D7 (kontroll: avvik ≈ 0)."),
    ("DiD vektet kontrast, tilnærming (log)", "=DiD_kontrast!O4", float(k_v), f"=Pipeline!B{pipe_row['D7 bestilt_post (WLS, n_mp)']}", "Pipelinen vekter i regresjonen (WLS med FE); enkel vektet middelverdi avviker litt. Forventet avvik."),
    ("DiD placebo uvektet kontrast (log)", "=DiD_placebo!O2", float(pl_uv), f"=Pipeline!B{pipe_row['D7 placebo (WLS)']}", "Pipelinens placebo er vektet og med felles profil per klasse; kontrasten her har egen profil per område. Sammenlignbar i størrelsesorden."),
    ("WCB-intervall D7, prosent (lav)", None, None, f"=(EXP(Pipeline!B{pipe_row['WCB KI lav (log)']})-1)*100", "Wild cluster bootstrap kan ikke regnes i Excel; kun omregning log → prosent."),
    ("WCB-intervall D7, prosent (høy)", None, None, f"=(EXP(Pipeline!B{pipe_row['WCB KI høy (log)']})-1)*100", ""),
    ("Forbruk H1 2026, husholdning+hytter, GWh", "=Forbruk!K15", float(h1.loc[2026, "sum"]), f"=Pipeline!B{pipe_row['Forbruk H1 2026, hush+hytter GWh']}", ""),
    ("Forbruksøkning H1 2026, prosent", "=Forbruk!M15", float((h1.loc[2026, "sum"] / h1.loc[2025, "sum"] - 1) * 100), None, ""),
    ("Forbruksandel bestillere", "=Aggregering!B4", float(andel_py), f"=Pipeline!B{pipe_row['Forbruksandel bestillere (post)']}", ""),
    ("Aggregert effekt, prosent (D7)", "=Aggregering!B7", None, None, "0,77 × 3,0 ≈ 2,3."),
    ("Dose sent, NO1 januar 2026", "=Dose!K6", float(dose_py["NO1"].get("2026-01", np.nan)), None, "Se Dose for alle måneder og områder."),
]
for i, (lab, ex, py_, pipe, kom) in enumerate(R, 2):
    put(wr, i, 1, lab)
    if ex: put(wr, i, 2, ex, GREEN, "0.000000")
    if py_ is not None: put(wr, i, 3, py_, BLUE, "0.000000")
    if pipe: put(wr, i, 4, pipe, GREEN, "0.000000")
    if ex and py_ is not None: put(wr, i, 5, f"=B{i}-C{i}", fmt="0.000000")
    if ex and pipe: put(wr, i, 6, f"=B{i}-D{i}", fmt="0.000000")
    put(wr, i, 7, kom, NOTE)
put(wr, len(R) + 3, 1, "Alle tall til godkjenning. Python-kontroll er regnet i s19 med samme kjede som Excel-formlene; Pipeline er statsmodels/linearmodels-resultatene i output/.", NOTE)

wb.move_sheet("Resultater", offset=-(len(wb.sheetnames) - 2))
wb.save(XLSX)
print(f"Skrevet {XLSX} | ark: {wb.sheetnames}")
print(f"Python-kontroll: M1 post {m1_b['post_np']:.6f} (pipeline {fe_m1.get('post_np'):.6f}) | M5 {m5_b['post_np']:.6f} (pipeline {fe_m5.get('post_np'):.6f}) | R2 M1 {m1_r2:.4f}")
